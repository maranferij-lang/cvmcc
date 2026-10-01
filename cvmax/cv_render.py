"""Оформлення зібраного CV: DOCX для завантаження і Markdown для перегляду."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .schemas import BuiltCV, CVEntry

FONTS_DIR = __import__("pathlib").Path(__file__).resolve().parents[1] / "assets" / "fonts"

PAGE_WIDTH_IN = 8.27  # A4
MARGIN_IN = 0.6
RIGHT_TAB = Inches(PAGE_WIDTH_IN - 2 * MARGIN_IN)


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "444444"}.items():
        bottom.set(qn(key), value)
    borders.append(bottom)
    # Word вимагає строгого порядку елементів у w:pPr: межі йдуть перед shd, tabs, spacing, ind, jc тощо.
    p_pr.insert_element_before(
        borders, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange",
    )


def _spacing(paragraph, before: float = 0, after: float = 0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def split_gpa(details: list[str]) -> tuple[str, list[str]]:
    """Короткий рядок з GPA виносимо праворуч у рядок ступеня, як у класичних CV."""
    for i, d in enumerate(details):
        if d.strip().upper().startswith("GPA") and len(d) <= 30:
            return d.strip(), details[:i] + details[i + 1 :]
    return "", details


def has_placeholders(cv: BuiltCV) -> list[str]:
    """Заповнювачі на кшталт [X] чи github.com/..., які юзер ще не замінив."""
    texts = cv.contact_line + [cv.summary]
    for group in (cv.experience, cv.projects, cv.activities):
        texts += [b for e in group for b in e.bullets]
    texts += [d for e in cv.education for d in e.details]
    return [t for t in texts if "[" in t or "..." in t]


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    _spacing(p, before=8, after=3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    _bottom_border(p)


def _line_with_date(doc, left: str, right: str, bold: bool = True, italic: bool = False) -> None:
    p = doc.add_paragraph()
    _spacing(p, before=3)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left)
    r.bold, r.italic = bold, italic
    if right:
        r2 = p.add_run("\t" + right)
        r2.italic = italic


def _bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        _spacing(p)


def _entries(doc, title: str, entries: list[CVEntry]) -> None:
    if not entries:
        return
    _heading(doc, title)
    for e in entries:
        org = ", ".join(x for x in (e.organization, e.location) if x.strip())
        if org:
            _line_with_date(doc, org, e.dates)
            p = doc.add_paragraph()
            _spacing(p)
            p.add_run(e.title).italic = True
        else:
            _line_with_date(doc, e.title, e.dates)
        _bullets(doc, e.bullets)


def render_docx(cv: BuiltCV) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(PAGE_WIDTH_IN), Inches(11.69)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(MARGIN_IN))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(name, after=2)
    run = name.add_run(cv.full_name)
    run.bold = True
    run.font.size = Pt(18)
    if cv.contact_line:
        contacts = doc.add_paragraph(" | ".join(cv.contact_line))
        contacts.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(contacts, after=4)
    if cv.summary.strip():
        _spacing(doc.add_paragraph(cv.summary.strip()), before=2)

    if cv.education:
        _heading(doc, "Education")
        for ed in cv.education:
            _line_with_date(doc, ", ".join(x for x in (ed.institution, ed.location) if x.strip()), ed.dates)
            gpa, details = split_gpa(ed.details)
            if ed.degree.strip() or gpa:
                _line_with_date(doc, ed.degree, gpa, bold=False, italic=True)
            _bullets(doc, details)
    _entries(doc, "Experience", cv.experience)
    _entries(doc, "Projects", cv.projects)
    _entries(doc, "Leadership & Activities", cv.activities)
    if cv.skills or cv.languages:
        _heading(doc, "Additional Information")
        for group in cv.skills:
            p = doc.add_paragraph()
            _spacing(p)
            p.add_run(f"{group.category}: ").bold = True
            p.add_run(", ".join(group.items))
        if cv.languages:
            p = doc.add_paragraph()
            _spacing(p)
            p.add_run("Languages: ").bold = True
            p.add_run("; ".join(cv.languages))
    if cv.awards:
        _heading(doc, "Awards")
        _bullets(doc, cv.awards)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_markdown(cv: BuiltCV) -> str:
    out = [f"## {cv.full_name}", " | ".join(cv.contact_line), ""]
    if cv.summary.strip():
        out += [cv.summary.strip(), ""]

    def entries(title: str, items: list[CVEntry]) -> None:
        if not items:
            return
        out.append(f"#### {title.upper()}")
        for e in items:
            out.append(f"**{e.title}**, {e.organization}{', ' + e.location if e.location else ''} · _{e.dates}_")
            out.extend(f"- {b}" for b in e.bullets)
            out.append("")

    if cv.education:
        out.append("#### EDUCATION")
        for ed in cv.education:
            out.append(f"**{ed.institution}**, {ed.degree} · _{ed.dates}_")
            out.extend(f"- {d}" for d in ed.details)
            out.append("")
    entries("Experience", cv.experience)
    entries("Projects", cv.projects)
    entries("Leadership & Activities", cv.activities)
    if cv.skills or cv.languages:
        out.append("#### ADDITIONAL INFORMATION")
        out.extend(f"- **{g.category}:** {', '.join(g.items)}" for g in cv.skills)
        if cv.languages:
            out.append(f"- **Languages:** {'; '.join(cv.languages)}")
        out.append("")
    if cv.awards:
        out.append("#### AWARDS")
        out.extend(f"- {a}" for a in cv.awards)
    return "\n".join(out)


# ---------------- PDF ----------------


def render_pdf(cv: BuiltCV) -> bytes:
    """Оформлене CV у PDF: A4, один стовпчик, без таблиць і картинок, текст виділяється (ATS)."""
    from fpdf import FPDF

    pdf = FPDF(format="A4", unit="mm")
    margin = 14
    pdf.set_margins(margin, 12, margin)
    pdf.set_auto_page_break(True, margin=12)
    family = "Liberation"
    for style, name in (("", "Regular"), ("B", "Bold"), ("I", "Italic"), ("BI", "BoldItalic")):
        pdf.add_font(family, style, str(FONTS_DIR / f"LiberationSans-{name}.ttf"))
    pdf.add_page()
    width = pdf.w - 2 * margin
    body = 10
    line = 4.6

    def heading(text: str) -> None:
        pdf.ln(2.2)
        pdf.set_font(family, "B", 10.5)
        pdf.cell(0, 5, text.upper(), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.set_draw_color(60, 60, 60)
        pdf.set_line_width(0.3)
        pdf.line(margin, y, margin + width, y)
        pdf.ln(1.2)

    def left_right(left: str, right: str, style_left: str = "B", style_right: str = "") -> None:
        pdf.set_font(family, style_right, body)
        right_w = pdf.get_string_width(right) + 1 if right else 0
        pdf.set_font(family, style_left, body)
        pdf.cell(width - right_w, line, left)
        if right:
            pdf.set_font(family, style_right, body)
            pdf.cell(right_w, line, right, align="R")
        pdf.ln(line)

    def bullets(items: list[str]) -> None:
        pdf.set_font(family, "", body)
        indent = 4
        for item in items:
            pdf.set_x(margin + 1)
            pdf.cell(indent - 1, line, "\u2022")
            pdf.multi_cell(width - indent, line, item.strip(), align="L", new_x="LMARGIN", new_y="NEXT")

    def entries(title: str, items: list[CVEntry]) -> None:
        if not items:
            return
        heading(title)
        for e in items:
            org = ", ".join(x for x in (e.organization, e.location) if x.strip())
            if org:
                left_right(org, e.dates, "B", "")
                pdf.set_font(family, "I", body)
                pdf.cell(0, line, e.title, new_x="LMARGIN", new_y="NEXT")
            else:
                left_right(e.title, e.dates, "B", "")
            bullets(e.bullets)
            pdf.ln(0.8)

    pdf.set_font(family, "B", 17)
    pdf.cell(0, 8, cv.full_name, align="C", new_x="LMARGIN", new_y="NEXT")
    if cv.contact_line:
        pdf.set_font(family, "", body)
        pdf.cell(0, line + 0.5, " | ".join(cv.contact_line), align="C", new_x="LMARGIN", new_y="NEXT")
    if cv.summary.strip():
        pdf.ln(1)
        pdf.set_font(family, "", body)
        pdf.multi_cell(0, line, cv.summary.strip(), align="L", new_x="LMARGIN", new_y="NEXT")

    if cv.education:
        heading("Education")
        for ed in cv.education:
            left_right(", ".join(x for x in (ed.institution, ed.location) if x.strip()), ed.dates)
            gpa, details = split_gpa(ed.details)
            if ed.degree.strip() or gpa:
                left_right(ed.degree, gpa, "I", "")
            bullets(details)
    entries("Experience", cv.experience)
    entries("Projects", cv.projects)
    entries("Leadership & Activities", cv.activities)
    if cv.skills or cv.languages:
        heading("Additional Information")
        rows = [(g.category, ", ".join(g.items)) for g in cv.skills]
        if cv.languages:
            rows.insert(len([r for r in rows if r[0].lower() != "interests"]), ("Languages", ", ".join(cv.languages)))
        for label, text in rows:
            pdf.set_font(family, "B", body)
            label_w = pdf.get_string_width(f"{label}: ") + 0.5
            pdf.cell(label_w, line, f"{label}: ")
            pdf.set_font(family, "", body)
            pdf.multi_cell(width - label_w, line, text, align="L", new_x="LMARGIN", new_y="NEXT")
    if cv.awards:
        heading("Awards")
        bullets(cv.awards)
    return bytes(pdf.output())
