"""Applying accepted "before / after" edits to the CV text, and export."""

from __future__ import annotations

import difflib
import io
import re
from dataclasses import dataclass, field

from docx import Document

from .schemas import Edit


@dataclass
class ApplyReport:
    text: str
    applied: list[Edit] = field(default_factory=list)
    added: list[Edit] = field(default_factory=list)
    not_found: list[Edit] = field(default_factory=list)


def _loose_pattern(snippet: str) -> re.Pattern[str]:
    # Text from a PDF often has extra line breaks and spaces, so we search with flexible gaps.
    words = snippet.split()
    return re.compile(r"\s+".join(re.escape(w) for w in words))


def _norm(text: str) -> str:
    return " ".join(text.lower().split())


def _fuzzy_span(text: str, snippet: str, threshold: float = 0.9) -> tuple[int, int] | None:
    """Finds a fragment almost identical to snippet (the model sometimes quotes with a small error: Build/Built).

    We compare snippet with windows of 1-6 neighboring lines and take the most similar one, if the similarity is at least threshold.
    """
    target = _norm(snippet)
    if len(target) < 20:
        return None
    lines = text.split("\n")
    offsets, pos = [], 0
    for line in lines:
        offsets.append(pos)
        pos += len(line) + 1
    best: tuple[float, int, int] | None = None
    for i in range(len(lines)):
        for j in range(i, min(i + 6, len(lines))):
            window = "\n".join(lines[i : j + 1])
            ratio = difflib.SequenceMatcher(None, _norm(window), target).ratio()
            if best is None or ratio > best[0]:
                start = offsets[i] + (len(lines[i]) - len(lines[i].lstrip(" •-\t")))
                best = (ratio, start, offsets[j] + len(lines[j]))
    if best and best[0] >= threshold:
        return best[1], best[2]
    return None


def apply_edits(cv_text: str, edits: list[Edit]) -> ApplyReport:
    report = ApplyReport(text=cv_text)
    for edit in edits:
        before = edit.before.strip()
        if not before:
            if edit.after.strip():
                report.added.append(edit)
            continue
        if before in report.text:
            report.text = report.text.replace(before, edit.after.strip(), 1)
            report.applied.append(edit)
            continue
        match = _loose_pattern(before).search(report.text)
        if match:
            report.text = report.text[: match.start()] + edit.after.strip() + report.text[match.end() :]
            report.applied.append(edit)
            continue
        span = _fuzzy_span(report.text, before)
        if span:
            start, end = span
            report.text = report.text[:start] + edit.after.strip() + report.text[end:]
            report.applied.append(edit)
        else:
            report.not_found.append(edit)

    # Removed items leave empty lines, we clean them up.
    report.text = re.sub(r"\n{3,}", "\n\n", report.text).strip()

    if report.added:
        by_section: dict[str, list[str]] = {}
        for e in report.added:
            by_section.setdefault(e.section.strip() or "Other", []).append(e.after.strip())
        extra = ["", "", "=== NEW ITEMS (move them into the right section) ==="]
        for section, items in by_section.items():
            extra.append(f"\n{section}")
            extra.extend(f"- {item}" for item in items)
        report.text += "\n".join(extra)
    return report


def changes_markdown(edits: list[Edit]) -> str:
    lines = ["# GetCVmax: accepted edits", ""]
    for i, e in enumerate(edits, 1):
        lines.append(f"## {i}. {e.section}")
        if e.before.strip():
            lines.append(f"**Before:** {e.before.strip()}")
        lines.append(f"**After:** {e.after.strip() or '(remove)'}")
        lines.append(f"_Why:_ {e.reason.strip()}")
        lines.append("")
    return "\n".join(lines)


def text_to_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- Check for invented facts ----------------

# Tools and technologies that the model most often adds on its own.
TECH_TERMS = {
    "python", "pandas", "numpy", "scipy", "matplotlib", "seaborn", "plotly", "scikit-learn", "sklearn",
    "pytorch", "tensorflow", "keras", "xgboost", "statsmodels", "beautifulsoup", "selenium", "scrapy",
    "sql", "postgresql", "mysql", "sqlite", "bigquery", "snowflake", "mongodb", "spark", "hadoop", "airflow",
    "excel", "vba", "vlookup", "xlookup", "power query", "powerquery", "pivot", "tableau", "power bi", "looker",
    "r", "stata", "spss", "eviews", "matlab", "sas", "jasp",
    "git", "github", "docker", "kubernetes", "aws", "azure", "gcp", "linux", "jira", "confluence", "notion",
    "figma", "miro", "canva", "hubspot", "salesforce", "sap", "1c", "crm", "a/b", "etl", "api",
    "javascript", "typescript", "react", "node", "java", "kotlin", "swift", "c++", "c#", "go", "django",
    "flask", "fastapi", "langchain", "llm", "gpt",
    # marketing and product analytics
    "google analytics", "ga4", "meta business suite", "ads manager", "google ads", "semrush", "ahrefs",
    "mailchimp", "hotjar", "amplitude", "mixpanel", "dbt", "seo", "ppc", "cac", "ltv", "roas", "ctr", "cpc", "cpa",
}
_TOKEN = re.compile(r"[A-Za-z][A-Za-z0-9+#./-]*[A-Za-z0-9+#]|[A-Za-z]|\d[\d.,]*")
_PHRASES = sorted((t for t in TECH_TERMS if " " in t), key=len, reverse=True)


def _suspicious(token: str) -> bool:
    low = token.lower()
    if low in TECH_TERMS:
        return True
    if any(ch.isdigit() for ch in token):
        return True
    if re.search(r"\.(com|org|io|dev|net|me|ua)\b|/", low):  # link: linkedin.com/in/...
        return True
    if len(token) >= 2 and token.isupper():  # SQL, VLOOKUP, KPI
        return True
    return any(c.isupper() for c in token[1:]) and any(c.islower() for c in token)  # PostgreSQL, NumPy


# Numbers that in a CV are often written as words.
_NUMBER_WORDS = {
    "half": "50", "quarter": "25", "third": "33", "twice": "2", "double": "2", "triple": "3",
    "one": "1", "two": "2", "three": "3", "four": "4", "five": "5", "six": "6", "seven": "7",
    "eight": "8", "nine": "9", "ten": "10", "dozen": "12", "hundred": "100", "thousand": "1000",
}


# Ukrainian numerals by stem: "п'ять" (five), "пятьма" (by five), "десяти" (of ten)...
_UA_NUMBER_STEMS = {
    "один": "1", "одн": "1", "два": "2", "дві": "2", "двох": "2", "три": "3", "трьох": "3",
    "чотир": "4", "п'ят": "5", "пят": "5", "шіст": "6", "шест": "6", "сім": "7", "сем": "7",
    "вісім": "8", "восьм": "8", "дев'ят": "9", "девят": "9", "десят": "10", "двадцят": "20",
    "тридцят": "30", "сорок": "40", "сто": "100", "тисяч": "1000", "половин": "50",
}

# Platform names that users write in Cyrillic.
_ALIASES = {
    "tiktok": ("тік ток", "тікток", "тик ток", "тикток"),
    "instagram": ("інстаграм", "инстаграм", "інста"),
    "youtube": ("ютуб", "ютюб"),
    "telegram": ("телеграм", "тг"),
    "linkedin": ("лінкедин", "лінкедін"),
    "claude": ("клод", "клоді", "клода"),
    "github": ("гітхаб", "гитхаб"),
    "ai": ("ші", "штучн", "llm", "gpt", "чатгпт", "chatgpt", "claude", "gemini"),
    "kse": ("кше", "київська школа економіки"),
    "sql": ("скл", "сікуель"),
    "excel": ("ексель", "эксель"),
}


def _numbers(text: str) -> set[str]:
    found = {re.sub(r"\D", "", n) for n in re.findall(r"\d[\d.,]*", text)}
    for word in re.findall(r"[a-z]+", text.lower()):
        if word in _NUMBER_WORDS:
            found.add(_NUMBER_WORDS[word])
    for word in re.findall(r"[а-яіїєґ']+", text.lower()):
        for stem, value in _UA_NUMBER_STEMS.items():
            if word.startswith(stem) and len(word) <= len(stem) + 4:
                found.add(value)
                break
    return found


def _known_alias(token: str, known_low: str) -> bool:
    # A synonym must stand at the start of a word: "ші" in "ШІ-школа" (AI school) counts, but in "інші" (others) it does not.
    return any(re.search(r"(?<![a-zа-яіїєґ])" + re.escape(alias), known_low)
               for alias in _ALIASES.get(token.lower(), ()))


def unverified_terms(after: str, known_text: str) -> list[str]:
    """Tool names and numbers from an edit that are in neither the CV nor the user's answers.

    This is a heuristic: it does not catch everything, but it highlights the model's most frequent inventions.
    """
    visible = re.sub(r"\[[^\]]*\]", " ", after)  # placeholders in brackets do not count
    known_low = known_text.lower()
    known_nums = _numbers(known_text)
    flagged: list[str] = []

    for phrase in _PHRASES:
        if phrase in visible.lower() and phrase not in known_low:
            flagged.append(phrase.title())
        visible = re.sub(re.escape(phrase), " ", visible, flags=re.I)

    for token in _TOKEN.findall(visible):
        token = token.strip(".,/-")
        if not token or not _suspicious(token):
            continue
        if token[0].isdigit():
            digits = re.sub(r"\D", "", token)
            if not digits or digits in known_nums or digits.rstrip("0") in known_nums:
                continue
        elif re.search(r"(?<![a-z0-9])" + re.escape(token.lower()) + r"(?![a-z0-9])", known_low):
            continue
        elif _known_alias(token, known_low):
            continue
        if token not in flagged:
            flagged.append(token)
    return flagged


_FACT_NUMBER = re.compile(r"\d[\d.,]*\s?%?")


def lost_facts(before: str, after: str) -> list[str]:
    """Numbers from the original item that disappeared in the rewritten one. An empty after (deletion) is not checked.

    A strong item stands on numbers, and the model sometimes loses them when rewriting. The user should see this.
    """
    if not before.strip() or not after.strip():
        return []
    after_digits = {re.sub(r"\D", "", n) for n in _FACT_NUMBER.findall(after)}
    lost = []
    for raw in _FACT_NUMBER.findall(before):
        number = raw.strip().rstrip(".,")
        if re.sub(r"\D", "", number) not in after_digits and number not in lost:
            lost.append(number)
    return lost



def changed_action(before: str, after: str) -> tuple[str, str] | None:
    """The main verb of the item changed to another one that was not in the original ("Produce" -> "Research").

    This is how the model sometimes passes off one experience as another. Returns (was, became) or None.
    """
    def first_word(text: str) -> str:
        m = re.match(r"[\s•▪\-–—*]*([A-Za-z]+)", text)
        return m.group(1) if m else ""

    old, new = first_word(before), first_word(after)
    if not old or not new or old[:4].lower() == new[:4].lower():
        return None
    if not old[0].isupper() or not new[0].isupper():
        return None  # not an item that starts with a verb
    if re.search(r"\b" + re.escape(new[:5].lower()), before.lower()):
        return None  # the new word is already in the original
    return old, new
