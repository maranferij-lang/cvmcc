import io
import json
from types import SimpleNamespace

import pytest
from docx import Document

from cvmax.builder import BuilderDraft, build_cv, builder_build_system, next_builder_question
from cvmax.cv_render import render_docx, render_markdown
from cvmax.db import DBError, MemoryDB, SupabaseDB
from cvmax.demo import FakeClient, demo_built_cv
from cvmax.grill import GrillSession, answer


class FakeHTTP:
    def __init__(self, status=200, body=None):
        self.calls, self.status, self.body = [], status, body

    def post(self, url, json=None, headers=None, timeout=None):
        self.calls.append((url, json, headers))
        content = b"" if self.body is None else b"x"
        return SimpleNamespace(status_code=self.status, content=content, text="err",
                               json=lambda: self.body)


def test_supabase_rpc_sends_token_and_key():
    http = FakeHTTP(body={"allowed": True})
    db = SupabaseDB("https://p.supabase.co/", "sb_publishable_x", "tok", session=http)
    assert db.consume("a@b.c", "analysis", 5, 60) == {"allowed": True}
    url, body, headers = http.calls[0]
    assert url == "https://p.supabase.co/rest/v1/rpc/cvmax_consume"
    assert body == {"p_token": "tok", "p_user_key": "a@b.c", "p_kind": "analysis",
                    "p_user_limit": 5, "p_global_limit": 60}
    assert headers["apikey"] == "sb_publishable_x" and "Authorization" not in headers


def test_supabase_legacy_jwt_key_adds_bearer_and_errors_raise():
    http = FakeHTTP(status=401, body={"message": "unauthorized"})
    db = SupabaseDB("https://p.supabase.co", "a.b.c", "bad", session=http)
    with pytest.raises(DBError):
        db.list_results("a@b.c")
    assert http.calls[0][2]["Authorization"] == "Bearer a.b.c"


def test_memory_db_limits_per_user_and_global():
    db = MemoryDB()
    assert db.consume("u1", "career", 2, 3)["allowed"]
    assert db.consume("u1", "career", 2, 3)["allowed"]
    assert db.consume("u1", "career", 2, 3) == {"allowed": False, "reason": "user", "used": 2, "limit": 2}
    assert db.consume("u2", "career", 2, 3)["allowed"]
    assert db.consume("u3", "career", 2, 3)["reason"] == "global"
    assert db.consume("u3", "analysis", 2, 3)["allowed"]  # інший вид не впливає


def draft():
    return BuilderDraft(full_name="Maya Chen", email="o@e.com", phone="", city="Kyiv", links="",
                        program="economics_big_data", status="2 курс", grad_year="2027", gpa="",
                        target_role="Data Analyst", notes="стажування в продажах, звіти в Excel",
                        feedback_language="Ukrainian")


def test_builder_flow_with_demo_client():
    client, g = FakeClient(), GrillSession(max_questions=2)
    q = next_builder_question(client, draft(), g)
    assert q is not None
    answer(g, "6 івентів на 300 студентів")
    cv = build_cv(client, draft(), g)
    assert cv.full_name == "Maya Chen" and g.finished
    sent = client.calls[-1]["messages"][0]["content"][0]["text"]
    assert "6 івентів" in sent and "стажування в продажах" in sent
    assert "Never invent" in builder_build_system(draft())


def test_render_docx_and_markdown():
    cv = demo_built_cv()
    doc = Document(io.BytesIO(render_docx(cv)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Maya Chen" in text and "EDUCATION" in text and "EXPERIENCE" in text
    assert "PROJECTS" not in text  # порожні розділи не друкуються
    md = render_markdown(cv)
    assert "#### EXPERIENCE" in md and "Sales Intern" in md
    json.dumps(cv.model_dump())  # результат можна зберегти в базу


def test_heading_border_comes_before_spacing_in_ppr():
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(render_docx(demo_built_cv())))
    for p in doc.paragraphs:
        ppr = p._p.pPr
        if ppr is None or ppr.find(qn("w:pBdr")) is None:
            continue
        tags = [child.tag for child in ppr]
        assert tags.index(qn("w:pBdr")) < tags.index(qn("w:spacing"))
