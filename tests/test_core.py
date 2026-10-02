import io

import pytest
from docx import Document

from cvmax import config
from cvmax.analyze import analyze_cv
from cvmax.cv_input import CVReadError, load_cv
from cvmax.demo import DEMO_CV_TEXT, FakeClient, demo_analysis
from cvmax.edits import apply_edits, changes_markdown, text_to_docx
from cvmax.grill import GrillSession, answer, finalize, next_question
from cvmax.profile import Profile
from cvmax.prompts import analysis_system, load_rubric
from cvmax.schemas import Edit


def make_profile(**over):
    base = dict(
        program="economics_big_data", status="3 курс", background="", target_role="Data Analyst",
        company_type="Fintech / bank", company_details="", level="Internship", region="EU",
        vacancy_text="", feedback_language="Ukrainian",
    )
    base.update(over)
    return Profile(**base)


def cv():
    return load_cv("cv.docx", _docx_bytes(DEMO_CV_TEXT))


def _docx_bytes(text):
    buf = io.BytesIO()
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(buf)
    return buf.getvalue()


def test_target_clarity_levels():
    assert make_profile(company_type="Any / not sure").target_clarity()[0] == "low"
    assert make_profile().target_clarity()[0] == "medium"
    assert make_profile(vacancy_text="x" * 300).target_clarity()[0] == "high"


def test_every_program_has_rubric():
    from cvmax.profile import PROGRAMS
    for key in PROGRAMS:
        if key != "other":
            assert key in load_rubric(key) or len(load_rubric(key)) > len(load_rubric("other"))


def test_prompt_mentions_language_and_rubric():
    p = analysis_system(make_profile(feedback_language="English"))
    assert "English" in p and "Role disambiguation" in p and "Economics and Big Data" in p


def test_load_docx_and_blocks():
    c = cv()
    assert "Responsible for making reports in Excel" in c.text
    assert c.as_content_blocks()[0]["type"] == "text"


def test_load_rejects_bad_types():
    with pytest.raises(CVReadError):
        load_cv("cv.doc", b"x")
    with pytest.raises(CVReadError):
        load_cv("cv.pdf", b"not a pdf")


def test_analyze_sends_expected_request():
    client = FakeClient()
    a = analyze_cv(client, make_profile(), cv())
    assert a.overall_score == 58
    call = client.calls[0]
    assert call["model"] == config.MODEL
    assert call["output_config"] == {"effort": config.EFFORT_ANALYSIS}
    assert call["extra_body"] == {"fallbacks": "default"}
    assert "<vacancy_text>" in call["messages"][0]["content"][-1]["text"]


def test_apply_edits_replace_remove_add_and_loose_match():
    text = "Responsible for making\n reports in Excel\nDate of birth: 01.01.2004\nEnd"
    report = apply_edits(text, demo_analysis().edits)
    assert "Built [N] weekly Excel reports" in report.text
    assert "Date of birth" not in report.text
    assert "SQL (joins" in report.text and "NEW ITEMS" in report.text
    assert not report.not_found


def test_apply_edits_reports_missing():
    e = Edit(section="X", before="nonexistent line", after="y", reason="r", priority="low")
    assert apply_edits("abc", [e]).not_found == [e]


def test_grill_flow_stops_and_finalizes():
    client, p, c = FakeClient(), make_profile(), cv()
    g = GrillSession()
    q = next_question(client, p, c, g)
    assert q and g.pending is q
    for _ in range(10):
        if g.pending is None:
            break
        answer(g, "I organised 6 events for 300 students")
        next_question(client, p, c, g)
    assert g.finished and len(g.turns) == 3
    assert "Q1:" in g.transcript()
    result = finalize(client, p, c, g)
    assert result.edits[0].before == "Member of Student Council"


def test_grill_respects_question_limit():
    client, g = FakeClient(), GrillSession(max_questions=1)
    next_question(client, make_profile(), cv(), g)
    answer(g, "")
    assert next_question(client, make_profile(), cv(), g) is None
    assert g.finished and g.turns[0].answer == "(skipped)"


def test_exports():
    md = changes_markdown(demo_analysis().edits)
    assert "**Before:**" in md and "(remove)" in md
    assert text_to_docx("a\nb")[:2] == b"PK"


def test_unverified_terms_flags_invented_tools_and_numbers():
    from cvmax.edits import unverified_terms
    known = "Skills: Excel, Python, SQL. Made 4 reports in 12 branches. Kyiv School of Economics"
    after = "Built 4 reports for 12 branches using Python (BeautifulSoup, Pandas), SQL (PostgreSQL), VLOOKUP and Power BI, saving 30%"
    flagged = unverified_terms(after, known)
    assert {"BeautifulSoup", "Pandas", "PostgreSQL", "VLOOKUP", "Power Bi", "30"} <= set(flagged)
    assert not {"Python", "SQL", "4", "12", "Built"} & set(flagged)


def test_unverified_terms_ignores_placeholders_and_known_facts():
    from cvmax.edits import unverified_terms
    assert unverified_terms("Used [Pandas?] to clean [X] rows in Excel", "Excel") == []
    assert unverified_terms("Analysed 20,000 listings", "20000 оголошень") == []
    assert unverified_terms("a | linkedin.com/in/andrii", "a") == ["linkedin.com/in/andrii"]


def test_unverified_terms_understands_number_words():
    from cvmax.edits import unverified_terms
    assert unverified_terms("Automated ~50% of daily posts", "roughly half of the posts") == []


def test_next_question_instruction_lists_covered_items():
    from cvmax.grill import QA, GrillSession, next_question_instruction
    g = GrillSession(turns=[QA("About the survey project?", "", "it was a side project")])
    text = next_question_instruction(g)
    assert "About the survey project?" in text and "Do NOT ask" in text
    assert "Do NOT ask" not in next_question_instruction(GrillSession())


def test_unverified_terms_understands_ukrainian_answers():
    from cvmax.edits import unverified_terms
    known = "вели тік ток та інстаграм клубу; провели 10 зустрічей, ще п'ять заплановано"
    assert unverified_terms("Ran TikTok and Instagram for the club; 10 meetups, 5 more planned", known) == []
    assert unverified_terms("Used Kubernetes", known) == ["Kubernetes"]


def test_apply_edits_tolerates_small_quote_mistakes():
    text = "EXPERIENCE\n• Built a Telegram bot that reminds club members about meetings\nand collects RSVPs in a sheet\nPROJECTS"
    e = Edit(section="Experience", before="Build a Telegram bot that reminds club members about meetings and collects RSVPs in a sheet",
             after="Built a bot that saves 1 hour daily", reason="r", priority="high")
    report = apply_edits(text, [e])
    assert report.applied == [e]
    assert "• Built a bot that saves 1 hour daily\nPROJECTS" in report.text
    unrelated = Edit(section="X", before="Completely different sentence about something else entirely",
                     after="y", reason="r", priority="low")
    assert apply_edits(text, [unrelated]).not_found == [unrelated]


def test_length_note_asks_for_cuts_only_when_over_one_page():
    from cvmax.analyze import length_note
    from cvmax.cv_input import CVFile
    short = CVFile("a.txt", "word " * 300)
    long = CVFile("b.txt", "word " * 900)
    two_pages = CVFile("c.pdf", "word " * 500, pages=2)
    assert "fits on one page" in length_note(short)
    assert "ONE page" in length_note(long) and "300 words" in length_note(long)
    assert "ONE page" in length_note(two_pages)
    assert length_note(CVFile("scan.pdf", "")) == ""


def test_grill_alternates_discover_and_deepen():
    from cvmax.grill import QA, GrillSession, next_question_instruction
    assert "DISCOVER" in next_question_instruction(GrillSession())
    one = GrillSession(turns=[QA("q", "", "a")])
    assert "DEEPEN" in next_question_instruction(one)


def test_cosmetic_edits_and_off_topic_gaps_are_dropped():
    from cvmax.analyze import drop_noise, is_cosmetic
    from cvmax.demo import demo_analysis
    from cvmax.schemas import Edit, Gap

    assert is_cosmetic("Sep 2023 – Jun 2024", "September 2023 - June 2024")
    assert is_cosmetic("Data Analyst, ACME | 2023-2024", "Data Analyst, Acme · 2023 – 2024")
    assert not is_cosmetic("Built 5 reports", "Built 12 weekly reports")
    assert not is_cosmetic("Date of birth: 01.01.2004", "")  # видалення не косметика

    a = demo_analysis()
    a.edits.append(Edit(section="Experience", before="Analyst, Acme, Sep 2023 - Jun 2024",
                        after="Analyst, Acme, September 2023 – June 2024", reason="date format", priority="low"))
    a.gaps.append(Gap(item="Networking", why_it_matters="x", how_to_close="Coffee chats with analysts",
                      time_estimate="ongoing", impact="medium"))
    a = drop_noise(a)
    assert all("September 2023" not in e.after for e in a.edits)
    assert all("Networking" != g.item for g in a.gaps)
    assert any("SQL" in g.item for g in a.gaps)  # корисні поради лишаються


def test_date_only_edits_are_dropped():
    from cvmax.analyze import is_cosmetic
    from cvmax.prompts import analysis_system
    from datetime import date

    assert is_cosmetic("Kyiv School of Economics Expected Graduation: June 2030",
                       "Kyiv School of Economics Expected Graduation: June [2028?]")
    assert is_cosmetic("Prozorro Defense Procurement Digest Sep 2026", "Prozorro Defense Procurement Digest [Sep 2024?]")
    assert not is_cosmetic("Economics Media Producer", "Operations & Data Analyst (Media Production)")
    assert str(date.today().year) in analysis_system(make_profile())


def test_changed_action_and_already_known_skill():
    from cvmax.analyze import drop_noise
    from cvmax.demo import demo_analysis
    from cvmax.edits import changed_action
    from cvmax.schemas import Gap

    assert changed_action("• Produce the outlet's podcast: source guests", "• Research industrial sectors") == ("Produce", "Research")
    assert changed_action("• Introduced automated clip production", "• Automated the clip pipeline") is None
    assert changed_action("Built 5 reports", "Built 12 weekly reports") is None

    a = demo_analysis()  # у правках є «SQL (joins...)», але без знака питання
    a.edits[1].after = "Python, [SQL?]"
    a.gaps = [Gap(item="SQL (joins)", why_it_matters="x", how_to_close="Take a course.", time_estimate="3 weeks", impact="high")]
    a = drop_noise(a)
    assert a.gaps[0].how_to_close.startswith("If you already use it")
